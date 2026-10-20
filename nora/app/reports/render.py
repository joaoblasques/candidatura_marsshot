"""Render a report's text into an Office file (PR.W / PR.X / PR.PPT).

The report body Nora drafts is plain email text; this turns it into the requested
Office document so the CEO gets both the inline report and a ready file. Pure and
local — no network, no external service.
"""

from __future__ import annotations

import io

from app.models import OutboundAttachment

# Format token -> (extension, MIME type). Mirrors the PR.<fmt> refinement tokens.
SUPPORTED_REPORT_FORMATS: dict[str, tuple[str, str]] = {
    "W": (".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "X": (".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "PPT": (".pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
}

_MAX_CHARS = 100_000  # bound the document we build from model output


def render_report(text: str, fmt: str) -> OutboundAttachment:
    """Render ``text`` into the file format named by ``fmt`` (a SUPPORTED key)."""
    if fmt not in SUPPORTED_REPORT_FORMATS:
        raise ValueError(f"unsupported report format: {fmt!r}")
    ext, content_type = SUPPORTED_REPORT_FORMATS[fmt]
    text = (text or "")[:_MAX_CHARS]
    builder = {"W": _docx, "X": _xlsx, "PPT": _pptx}[fmt]
    return OutboundAttachment(
        filename=f"Nora-report{ext}", content_type=content_type, content=builder(text)
    )


def _paragraphs(text: str) -> list[str]:
    """Split into non-empty blocks on blank lines (a simple, robust structure)."""
    blocks = [b.strip() for b in text.split("\n\n")]
    return [b for b in blocks if b] or ["(empty report)"]


def _docx(text: str) -> bytes:
    import docx

    doc = docx.Document()
    for para in _paragraphs(text):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx(text: str) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for i, line in enumerate((text or "").splitlines() or ["(empty report)"], start=1):
        ws.cell(row=i, column=1, value=line[:32_000])  # Excel per-cell limit guard
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx(text: str) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    blocks = _paragraphs(text)
    title = blocks[0]
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title[:200]
    for block in blocks[1:] or [title]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = block.splitlines()[0][:200]
        body = slide.placeholders[1].text_frame
        body.text = block[:3_000]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
