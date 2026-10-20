"""Local, exception-safe text extraction from forwarded attachments.

Only extracted TEXT ever leaves this module — raw files are never sent to the
LLM. Every path degrades to an ``unreadable`` result rather than raising, so a
hostile or corrupt file can never break the inbound pipeline.
"""

from __future__ import annotations

import io
import zipfile

from app.models import Attachment, ExtractedAttachment

PER_FILE_BYTES = 5_000_000
PER_FILE_TEXT_CHARS = 20_000
TOTAL_TEXT_CHARS = 50_000
MAX_UNCOMPRESSED_BYTES = 50_000_000  # zip-bomb guard for .docx/.xlsx

_TEXT_EXTS = {".txt", ".md", ".csv", ".log"}


def extract_attachments(attachments: list[Attachment]) -> list[ExtractedAttachment]:
    """Extract bounded text from each attachment, in order, honoring the total cap."""
    results: list[ExtractedAttachment] = []
    remaining = TOTAL_TEXT_CHARS
    for att in attachments:
        res = _extract_one(att, remaining)
        if res.status == "read":
            remaining -= len(res.text)
        results.append(res)
    return results


def _extract_one(att: Attachment, remaining: int) -> ExtractedAttachment:
    name = att.filename or "attachment"
    if att.content is None:
        reason = "too_large" if (att.size or 0) > PER_FILE_BYTES else "no_text_found"
        return ExtractedAttachment(filename=name, status="unreadable", reason=reason)
    if (att.size or len(att.content)) > PER_FILE_BYTES:
        return ExtractedAttachment(filename=name, status="unreadable", reason="too_large")
    try:
        raw = _dispatch(_ext(name), att.content)
    except Exception:
        return ExtractedAttachment(filename=name, status="unreadable", reason="extract_failed")
    if raw is None:
        return ExtractedAttachment(filename=name, status="unreadable", reason="unsupported_type")
    text = raw.strip()
    if not text:
        return ExtractedAttachment(filename=name, status="unreadable", reason="no_text_found")
    cap = min(PER_FILE_TEXT_CHARS, max(0, remaining))
    if cap <= 0:
        # The total text budget is spent — surface this file as NOT read (named in the
        # manifest) rather than a silent empty "read" block (transparency contract).
        return ExtractedAttachment(filename=name, status="unreadable", reason="budget_exhausted")
    return ExtractedAttachment(
        filename=name, status="read", text=text[:cap], truncated=len(text) > cap
    )


def _ext(name: str) -> str:
    name = name.lower()
    dot = name.rfind(".")
    return name[dot:] if dot >= 0 else ""


def _dispatch(ext: str, content: bytes) -> str | None:
    if ext == ".pdf":
        return _pdf(content)
    if ext == ".docx":
        return _docx(content)
    if ext == ".xlsx":
        return _xlsx(content)
    if ext in _TEXT_EXTS:
        return content.decode("utf-8", errors="replace")
    return None


def _guard_zip(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        if sum(i.file_size for i in zf.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError("decompressed size exceeds limit")


def _pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(content: bytes) -> str:
    _guard_zip(content)
    import docx

    doc = docx.Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx(content: bytes) -> str:
    _guard_zip(content)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)
